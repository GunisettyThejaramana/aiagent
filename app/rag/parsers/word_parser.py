from pathlib import Path
from docx import Document


class WordParser:
    """
    Extracts text from Microsoft Word (.docx) documents.
    """

    def parse(self, file_path: Path) -> str:
        """
        Read a Word document and return its text.

        Args:
            file_path (Path): Path to the Word document.

        Returns:
            str: Extracted text.
        """

        text = []

        try:
            document = Document(file_path)

            # Read paragraphs
            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    text.append(paragraph.text.strip())

            # Read tables
            for table in document.tables:
                for row in table.rows:
                    values = [
                        cell.text.strip()
                        for cell in row.cells
                    ]
                    text.append(" | ".join(values))

        except Exception as e:
            print(f"[WORD ERROR] {file_path}: {e}")

        return "\n".join(text)