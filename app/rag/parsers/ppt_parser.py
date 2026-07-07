from pathlib import Path
from pptx import Presentation


class PPTParser:
    """
    Extracts text from Microsoft PowerPoint (.pptx) files.
    """

    def parse(self, file_path: Path) -> str:
        """
        Read a PowerPoint presentation and return all text.

        Args:
            file_path (Path): Path to the PowerPoint file.

        Returns:
            str: Extracted text.
        """

        text = []

        try:
            presentation = Presentation(file_path)

            for slide_number, slide in enumerate(
                presentation.slides,
                start=1
            ):
                text.append(
                    f"\n===== Slide {slide_number} =====\n"
                )

                for shape in slide.shapes:

                    if hasattr(shape, "text"):

                        shape_text = shape.text.strip()

                        if shape_text:
                            text.append(shape_text)

        except Exception as e:
            print(f"[PPT ERROR] {file_path}: {e}")

        return "\n".join(text)