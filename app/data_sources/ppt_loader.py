from pathlib import Path

from pptx import Presentation
from langchain.schema import Document


def load_ppt(path: str):
    """
    Load a PowerPoint (.pptx) file and return LangChain Document objects.
    """

    ppt_path = Path(path)

    if not ppt_path.exists():
        print(f"PowerPoint file not found: {ppt_path}")
        return []

    try:
        presentation = Presentation(str(ppt_path))

        slides_text = []

        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_content = []

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_content.append(shape.text.strip())

            if slide_content:
                slides_text.append(
                    f"Slide {slide_number}\n" +
                    "\n".join(slide_content)
                )

        content = "\n\n".join(slides_text)

        return [
            Document(
                page_content=content,
                metadata={
                    "source": str(ppt_path),
                    "file_type": "pptx"
                }
            )
        ]

    except Exception as e:
        print(f"Error loading PowerPoint '{ppt_path}': {e}")
        return []